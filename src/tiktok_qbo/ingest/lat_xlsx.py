from pathlib import Path
from decimal import Decimal
from openpyxl import load_workbook
from tiktok_qbo.models import (
    NormalizedRow, StatementRow, PaymentRow, ReserveRow,
)
from tiktok_qbo.ingest.classify import classify
from tiktok_qbo.dates import parse_lat_date
from tiktok_qbo.money import to_money


def _sheet_rows(path: Path, sheet_name: str):
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Missing sheet: {sheet_name}")
    ws = wb[sheet_name]
    iter_rows = ws.iter_rows(values_only=True)
    header = next(iter_rows)
    for row in iter_rows:
        if row is None or all(v is None for v in row):
            continue
        yield dict(zip(header, row))
    wb.close()


def _require_shop(raw: dict, expected: str, path: Path):
    got = str(raw.get("Shop ID", "")).strip()
    if got and got != expected:
        raise ValueError(
            f"shop mismatch in {path.name}: expected {expected!r}, got {got!r}"
        )


def read_order_details(path: Path, shop_id: str) -> list[NormalizedRow]:
    out: list[NormalizedRow] = []
    for raw in _sheet_rows(Path(path), "Order details"):
        _require_shop(raw, shop_id, Path(path))
        classification = classify(raw)
        out.append(NormalizedRow(
            shop_id=shop_id,
            order_id=str(raw.get("Order ID", "") or ""),
            sku_id=str(raw.get("SKU ID", "") or ""),
            statement_id=str(raw.get("Statement ID", "") or ""),
            payment_id=str(raw.get("Payment ID", "") or ""),
            statement_date=parse_lat_date(raw.get("Statement date")),
            order_created_date=parse_lat_date(raw.get("Order created date")),
            order_shipment_date=parse_lat_date(raw.get("Order shipment date")),
            order_delivery_date=parse_lat_date(raw.get("Order delivery date")),
            row_type=str(raw.get("Type", "") or ""),
            classification=classification,
            customer_payment=to_money(raw.get("Customer payment")),
            customer_refund=to_money(raw.get("Customer refund")),
            gross_sales=to_money(raw.get("Gross sales")),
            quantity=int(raw.get("Quantity") or 0) if str(raw.get("Quantity") or "").strip() not in ("/", "", "None") else 0,
            product_name=str(raw.get("Product name", "") or ""),
            raw={str(k): (str(v) if v is not None else "") for k, v in raw.items() if k is not None},
        ))
    return out


def read_statements(path: Path, shop_id: str) -> list[StatementRow]:
    out: list[StatementRow] = []
    for raw in _sheet_rows(Path(path), "Statements"):
        _require_shop(raw, shop_id, Path(path))
        out.append(StatementRow(
            shop_id=shop_id,
            statement_id=str(raw.get("Statement ID", "") or ""),
            payment_id=str(raw.get("Payment ID", "") or ""),
            statement_date=parse_lat_date(raw.get("Statement date")),
            status=str(raw.get("Status", "") or ""),
            total_settlement_amount=to_money(raw.get("Total settlement amount")),
            net_sales=to_money(raw.get("Net sales")),
            shipping=to_money(raw.get("Shipping")),
            fees=to_money(raw.get("Fees")),
            adjustments=to_money(raw.get("Adjustments")),
            reserve_amount=to_money(raw.get("Reserve amount")),
            payable_amount=to_money(raw.get("Payable amount")),
        ))
    return out


def read_payments(path: Path, shop_id: str) -> list[PaymentRow]:
    out: list[PaymentRow] = []
    for raw in _sheet_rows(Path(path), "Payments"):
        _require_shop(raw, shop_id, Path(path))
        out.append(PaymentRow(
            shop_id=shop_id,
            payment_id=str(raw.get("Payment ID", "") or ""),
            payment_amount=to_money(raw.get("Payment amount")),
            payment_initiation_date=parse_lat_date(raw.get("Payment initiation date")),
            payment_completion_date=parse_lat_date(raw.get("Payment completion date")),
            bank_account_masked=str(raw.get("Bank account", "") or ""),
            status=str(raw.get("Status", "") or ""),
        ))
    return out


def read_reserves(path: Path, shop_id: str) -> list[ReserveRow]:
    out: list[ReserveRow] = []
    for raw in _sheet_rows(Path(path), "Reserve details"):
        _require_shop(raw, shop_id, Path(path))
        out.append(ReserveRow(
            shop_id=shop_id,
            statement_id=str(raw.get("Statement ID", "") or ""),
            reserve_id=str(raw.get("Reserve ID", "") or ""),
            reserve_amount=to_money(raw.get("Reserve amount")),
            reserve_date=parse_lat_date(raw.get("Reserve date")),
            release_date=parse_lat_date(raw.get("Release date")),
            status=str(raw.get("Status", "") or ""),
        ))
    return out
