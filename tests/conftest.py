# Fixtures added per-task.
from decimal import Decimal
from datetime import date
from pathlib import Path
import pytest
from openpyxl import Workbook

@pytest.fixture
def tiny_xlsx(tmp_path: Path) -> Path:
    """Build a minimal LAT xlsx covering sale + refund + chargeback + fee_only +
    one Statement + one Payment + one Reserve row."""
    wb = Workbook()

    od = wb.active
    od.title = "Order details"
    od.append([
        "Statement ID", "Payment ID", "Statement date", "Shop ID", "Order ID",
        "SKU ID", "Product name", "Quantity", "Type",
        "Order created date", "Order shipment date", "Order delivery date",
        "Gross sales", "Gross sales refund", "Customer payment", "Customer refund",
    ])
    # sale
    od.append(["S1","P1","2024-05-12","PLELNU","O1","SKU-A","Widget",1,"Order",
               "2024-04-28","2024-04-29","2024-04-30","15.99","0","15.99","0"])
    # refund
    od.append(["S1","P1","2024-05-12","PLELNU","O2","SKU-B","Gadget",1,"Order",
               "2024-04-10","2024-04-11","2024-04-12","0","-9.99","0","9.99"])
    # chargeback
    od.append(["S1","P1","2024-05-12","PLELNU","O3","SKU-C","Thing",1,"Chargeback",
               "2024-03-20","2024-03-21","2024-03-22","0","0","0","0"])
    # fee-only (platform fee row)
    od.append(["S1","P1","2024-05-12","PLELNU","O4","SKU-A","Widget",1,"Order",
               "2024-04-28","2024-04-29","2024-04-30","0","0","0","0"])

    s = wb.create_sheet("Statements")
    s.append(["Shop ID","Statement ID","Payment ID","Statement date","Status",
              "Total settlement amount","Net sales","Shipping","Fees",
              "Adjustments","Reserve amount","Payable amount"])
    # payable = 15.99 - 9.99 + 0 (fees) + 0 (shipping) + 0 (adj) - 0 (reserve) = 6.00
    s.append(["PLELNU","S1","P1","2024-05-12","Paid",
              "6.00","6.00","0","0","0","0","6.00"])

    p = wb.create_sheet("Payments")
    p.append(["Shop ID","Payment ID","Payment amount","Payment initiation date",
              "Payment completion date","Bank account","Status"])
    p.append(["PLELNU","P1","6.00","2024-05-12","2024-05-13","********9247","Completed"])

    r = wb.create_sheet("Reserve details")
    r.append(["Shop ID","Statement ID","Reserve ID","Reserve amount",
              "Reserve date","Release date","Status"])
    r.append(["PLELNU","S1","R1","0","2024-05-12","","Released"])

    path = tmp_path / "tiny.xlsx"
    wb.save(path)
    return path
