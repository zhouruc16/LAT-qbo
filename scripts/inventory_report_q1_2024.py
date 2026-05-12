"""Build a per-statement inventory report from a TikTok xlsx export.

Output: one sheet per Statement ID. Each sheet has a table of products sold
on that statement, grouped by (statement_id, SKU):

    SKU ID | Quantity | Product name | Unit cost | Line cost

Unit cost is joined from Master Table_updated.xlsx (falls back to
Master Table.xlsx if the _updated file is not present). SKUs without a
price in the master table get a blank Unit cost and Line cost.

Quantity = net sold on that statement = sale-row qty − refund-row qty.

Usage:
    python scripts/inventory_report_q1_2024.py
"""
from __future__ import annotations

from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tiktok_qbo.ingest.lat_xlsx import read_order_details, read_statements

SHOP = "USLCPLELNU"
SRC = Path.home() / "Downloads" / "1-3-2024 (1).xlsx"
OUT = Path.home() / "Downloads" / "1-3-2024_inventory_report.xlsx"
MASTER_UPDATED = Path.home() / "Downloads" / "Master Table_updated.xlsx"
MASTER_ORIG = Path.home() / "Downloads" / "Master Table.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="2F5597")  # blue header
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="DDDDDD")
TOTAL_FONT = Font(bold=True)
THIN_BORDER_FILL = PatternFill("solid", fgColor="F2F2F2")


def load_master_prices(path: Path) -> dict[str, tuple[str, Decimal | None]]:
    """Return {sku_id -> (product_name, unit_cost_or_None)} from a master xlsx."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    out: dict[str, tuple[str, Decimal | None]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        sku = str(row[0]).strip()
        if not sku.isdigit():
            continue
        name = str(row[1]).strip() if row[1] else ""
        price = None
        if len(row) >= 3 and row[2] not in (None, ""):
            try:
                price = Decimal(str(row[2]))
            except Exception:
                price = None
        out[sku] = (name, price)
    return out


def main() -> None:
    master_path = MASTER_UPDATED if MASTER_UPDATED.exists() else MASTER_ORIG
    print(f"Reading master: {master_path.name}")
    master = load_master_prices(master_path)
    priced = sum(1 for _n, p in master.values() if p is not None)
    print(f"  {len(master)} SKUs in master, {priced} have prices")
    print()

    rows = read_order_details(SRC, SHOP)
    statements = read_statements(SRC, SHOP)
    stmt_meta = {s.statement_id: s for s in statements}

    # group by (statement_id, sku_id)
    sale_qty: dict[tuple[str, str], int] = defaultdict(int)
    refund_qty: dict[tuple[str, str], int] = defaultdict(int)
    name_votes: dict[tuple[str, str], Counter] = defaultdict(Counter)

    for r in rows:
        if not r.sku_id or r.quantity is None:
            continue
        sku = str(r.sku_id).strip()
        if not sku.isdigit():
            continue
        key = (r.statement_id, sku)
        if r.product_name:
            name_votes[key][r.product_name] += r.quantity or 1
        if r.classification == "sale":
            sale_qty[key] += r.quantity
        elif r.classification == "refund":
            refund_qty[key] += r.quantity

    all_keys = set(sale_qty) | set(refund_qty)
    by_stmt: dict[str, list[tuple[str, int, str, Decimal | None]]] = defaultdict(list)
    for stmt_id, sku in all_keys:
        sold = sale_qty.get((stmt_id, sku), 0)
        refunded = refund_qty.get((stmt_id, sku), 0)
        net = sold - refunded
        if net == 0:
            continue
        # prefer name from master; else most-voted name from the xlsx rows
        master_name, master_price = master.get(sku, ("", None))
        if master_name:
            product_name = master_name
        else:
            nc = name_votes.get((stmt_id, sku), Counter())
            product_name = nc.most_common(1)[0][0] if nc else ""
        by_stmt[stmt_id].append((sku, net, product_name, master_price))

    # sort each statement's rows by quantity desc, then sku
    for stmt_id in by_stmt:
        by_stmt[stmt_id].sort(key=lambda x: (-x[1], x[0]))

    wb = Workbook()
    wb.remove(wb.active)

    sorted_stmts = sorted(
        by_stmt.keys(),
        key=lambda sid: (
            getattr(stmt_meta.get(sid), "statement_date", None) or "",
            sid,
        ),
    )

    summary_rows = []
    for stmt_id in sorted_stmts:
        meta = stmt_meta.get(stmt_id)
        date = getattr(meta, "statement_date", None)
        sheet_name = f"{date.isoformat()}_{stmt_id[-4:]}" if date else stmt_id[-12:]
        sheet_name = sheet_name[:31]
        ws = wb.create_sheet(title=sheet_name)

        # header rows
        ws.cell(row=1, column=1, value="Statement ID").font = TOTAL_FONT
        ws.cell(row=1, column=2, value=stmt_id)
        ws.cell(row=2, column=1, value="Statement Date").font = TOTAL_FONT
        ws.cell(row=2, column=2, value=date.isoformat() if date else "")

        header_row = 4
        headers = ["SKU ID", "Quantity", "Product name", "Unit cost", "Line cost"]
        for col_idx, h in enumerate(headers, start=1):
            c = ws.cell(row=header_row, column=col_idx, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")

        total_qty = 0
        total_line_cost = Decimal("0")
        priced_lines = 0
        unpriced_lines = 0
        last_r = header_row
        for i, (sku, qty, name, unit_cost) in enumerate(by_stmt[stmt_id], start=header_row + 1):
            ws.cell(row=i, column=1, value=sku)
            ws.cell(row=i, column=2, value=qty).alignment = Alignment(horizontal="center")
            ws.cell(row=i, column=3, value=name)
            if unit_cost is not None:
                uc_cell = ws.cell(row=i, column=4, value=float(unit_cost))
                uc_cell.number_format = '"$"#,##0.00'
                line_cost = (Decimal(qty) * unit_cost).quantize(Decimal("0.01"))
                lc_cell = ws.cell(row=i, column=5, value=float(line_cost))
                lc_cell.number_format = "#,##0.00"
                total_line_cost += line_cost
                priced_lines += 1
            else:
                # leave blank to mirror master-table treatment
                ws.cell(row=i, column=4, value=None)
                ws.cell(row=i, column=5, value=None)
                unpriced_lines += 1
            total_qty += qty
            last_r = i

        # TOTAL row
        total_row = last_r + 1
        ws.cell(row=total_row, column=1, value="TOTAL").font = TOTAL_FONT
        ws.cell(row=total_row, column=2, value=total_qty).font = TOTAL_FONT
        ws.cell(row=total_row, column=2).alignment = Alignment(horizontal="center")
        tc = ws.cell(row=total_row, column=5, value=float(total_line_cost) if total_line_cost else None)
        tc.font = TOTAL_FONT
        tc.number_format = "#,##0.00"
        for col in range(1, 6):
            ws.cell(row=total_row, column=col).fill = TOTAL_FILL

        widths = [22, 10, 70, 12, 12]
        for col_idx, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = w

        summary_rows.append(
            (
                date.isoformat() if date else "",
                stmt_id,
                len(by_stmt[stmt_id]),
                total_qty,
                priced_lines,
                unpriced_lines,
                float(total_line_cost),
            )
        )

    # summary sheet
    summary = wb.create_sheet(title="Summary", index=0)
    summary.cell(row=1, column=1, value="TikTok LELNU Q1 2024 Inventory Report").font = Font(
        bold=True, size=14
    )
    summary.cell(row=2, column=1, value=f"Source: {SRC.name}")
    summary.cell(row=3, column=1, value=f"Master: {master_path.name}")
    summary.cell(row=4, column=1, value=f"Statements: {len(by_stmt)}")
    headers = [
        "Statement Date",
        "Statement ID",
        "Unique SKUs",
        "Total Qty",
        "Priced Lines",
        "Unpriced Lines",
        "Total Line Cost",
    ]
    for col_idx, h in enumerate(headers, start=1):
        c = summary.cell(row=6, column=col_idx, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
    for i, srow in enumerate(summary_rows, start=7):
        for col_idx, v in enumerate(srow, start=1):
            cell = summary.cell(row=i, column=col_idx, value=v)
            if col_idx == 7:
                cell.number_format = "#,##0.00"
    widths = [14, 22, 12, 12, 14, 16, 18]
    for col_idx, w in enumerate(widths, start=1):
        summary.column_dimensions[get_column_letter(col_idx)].width = w

    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"  {len(by_stmt)} statements")
    print(f"  {sum(len(v) for v in by_stmt.values())} (statement, sku) rows")
    print(f"  Total priced lines: {sum(r[4] for r in summary_rows)}")
    print(f"  Total unpriced lines: {sum(r[5] for r in summary_rows)}")
    print(f"  Total line cost (priced only): ${sum(r[6] for r in summary_rows):,.2f}")


if __name__ == "__main__":
    main()
