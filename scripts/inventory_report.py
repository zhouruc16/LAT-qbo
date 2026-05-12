"""Build a per-statement inventory report from a TikTok xlsx export.

Output: one xlsx with one sheet per Statement ID. Each sheet has:

    SKU ID | Quantity | Product name | Unit cost | Line cost

Unit cost is joined from Master Table_updated.xlsx (falls back to
Master Table.xlsx if _updated is not present). SKUs without a price get
blank Unit cost / Line cost.

Quantity = sale-row qty − refund-row qty (net inventory movement).
Rows where net qty == 0 are dropped.

Usage:
    python scripts/inventory_report.py 1-3-2024
    python scripts/inventory_report.py 4-6-2024
    python scripts/inventory_report.py 7-9-2024

Source xlsx is looked up in ~/Downloads/. If `<stem>.xlsx` does not exist,
falls back to `<stem> (1).xlsx`.
"""
from __future__ import annotations

import sys
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tiktok_qbo.ingest.lat_xlsx import read_order_details, read_statements

SHOP = "USLCPLELNU"
DOWNLOADS = Path.home() / "Downloads"
MASTER_UPDATED = DOWNLOADS / "Master Table_updated.xlsx"
MASTER_ORIG = DOWNLOADS / "Master Table.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="DDDDDD")
TOTAL_FONT = Font(bold=True)


def resolve_source(stem: str) -> Path:
    primary = DOWNLOADS / f"{stem}.xlsx"
    if primary.exists():
        return primary
    fallback = DOWNLOADS / f"{stem} (1).xlsx"
    if fallback.exists():
        return fallback
    raise FileNotFoundError(f"Neither {primary.name} nor {fallback.name} exists in Downloads")


def load_master_prices(path: Path) -> dict[str, tuple[str, Decimal | None]]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    out: dict[str, tuple[str, Decimal | None]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        sku = str(row[0]).strip()
        if not sku.isdigit():
            continue
        name = str(row[1]).strip() if row[1] else ""
        price = None
        if len(row) >= 3 and row[2] not in (None, ""):
            try:
                price = Decimal(str(row[2]))
            except Exception:
                price = None
        out[sku] = (name, price)
    return out


def build_report(src: Path, out: Path, master: dict[str, tuple[str, Decimal | None]],
                 quarter_label: str) -> None:
    rows = read_order_details(src, SHOP)
    statements = read_statements(src, SHOP)
    stmt_meta = {s.statement_id: s for s in statements}

    sale_qty: dict[tuple[str, str], int] = defaultdict(int)
    refund_qty: dict[tuple[str, str], int] = defaultdict(int)
    name_votes: dict[tuple[str, str], Counter] = defaultdict(Counter)

    for r in rows:
        if not r.sku_id or r.quantity is None:
            continue
        sku = str(r.sku_id).strip()
        if not sku.isdigit():
            continue
        key = (r.statement_id, sku)
        if r.product_name:
            name_votes[key][r.product_name] += r.quantity or 1
        if r.classification == "sale":
            sale_qty[key] += r.quantity
        elif r.classification == "refund":
            refund_qty[key] += r.quantity

    by_stmt: dict[str, list[tuple[str, int, str, Decimal | None]]] = defaultdict(list)
    for stmt_id, sku in set(sale_qty) | set(refund_qty):
        sold = sale_qty.get((stmt_id, sku), 0)
        refunded = refund_qty.get((stmt_id, sku), 0)
        net = sold - refunded
        if net == 0:
            continue
        master_name, master_price = master.get(sku, ("", None))
        if master_name:
            product_name = master_name
        else:
            nc = name_votes.get((stmt_id, sku), Counter())
            product_name = nc.most_common(1)[0][0] if nc else ""
        by_stmt[stmt_id].append((sku, net, product_name, master_price))

    for stmt_id in by_stmt:
        by_stmt[stmt_id].sort(key=lambda x: (-x[1], x[0]))

    sorted_stmts = sorted(
        by_stmt.keys(),
        key=lambda sid: (
            getattr(stmt_meta.get(sid), "statement_date", None) or "",
            sid,
        ),
    )

    wb = Workbook()
    wb.remove(wb.active)
    summary_rows = []

    for stmt_id in sorted_stmts:
        meta = stmt_meta.get(stmt_id)
        date = getattr(meta, "statement_date", None)
        sheet_name = f"{date.isoformat()}_{stmt_id[-4:]}" if date else stmt_id[-12:]
        sheet_name = sheet_name[:31]
        ws = wb.create_sheet(title=sheet_name)

        ws.cell(row=1, column=1, value="Statement ID").font = TOTAL_FONT
        ws.cell(row=1, column=2, value=stmt_id)
        ws.cell(row=2, column=1, value="Statement Date").font = TOTAL_FONT
        ws.cell(row=2, column=2, value=date.isoformat() if date else "")

        header_row = 4
        headers = ["SKU ID", "Quantity", "Product name", "Unit cost", "Line cost"]
        for col_idx, h in enumerate(headers, start=1):
            c = ws.cell(row=header_row, column=col_idx, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")

        total_qty = 0
        total_line_cost = Decimal("0")
        priced_lines = 0
        unpriced_lines = 0
        last_r = header_row
        for i, (sku, qty, name, unit_cost) in enumerate(by_stmt[stmt_id], start=header_row + 1):
            ws.cell(row=i, column=1, value=sku)
            ws.cell(row=i, column=2, value=qty).alignment = Alignment(horizontal="center")
            ws.cell(row=i, column=3, value=name)
            if unit_cost is not None:
                uc = ws.cell(row=i, column=4, value=float(unit_cost))
                uc.number_format = '"$"#,##0.00'
                line_cost = (Decimal(qty) * unit_cost).quantize(Decimal("0.01"))
                lc = ws.cell(row=i, column=5, value=float(line_cost))
                lc.number_format = "#,##0.00"
                total_line_cost += line_cost
                priced_lines += 1
            else:
                unpriced_lines += 1
            total_qty += qty
            last_r = i

        total_row = last_r + 1
        ws.cell(row=total_row, column=1, value="TOTAL").font = TOTAL_FONT
        tq = ws.cell(row=total_row, column=2, value=total_qty)
        tq.font = TOTAL_FONT
        tq.alignment = Alignment(horizontal="center")
        tc = ws.cell(row=total_row, column=5, value=float(total_line_cost) if total_line_cost else None)
        tc.font = TOTAL_FONT
        tc.number_format = "#,##0.00"
        for col in range(1, 6):
            ws.cell(row=total_row, column=col).fill = TOTAL_FILL

        widths = [22, 10, 70, 12, 12]
        for col_idx, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = w

        summary_rows.append(
            (
                date.isoformat() if date else "",
                stmt_id,
                len(by_stmt[stmt_id]),
                total_qty,
                priced_lines,
                unpriced_lines,
                float(total_line_cost),
            )
        )

    summary = wb.create_sheet(title="Summary", index=0)
    summary.cell(row=1, column=1, value=f"TikTok LELNU {quarter_label} Inventory Report").font = Font(
        bold=True, size=14
    )
    summary.cell(row=2, column=1, value=f"Source: {src.name}")
    summary.cell(row=3, column=1, value=f"Statements: {len(by_stmt)}")
    headers = [
        "Statement Date",
        "Statement ID",
        "Unique SKUs",
        "Total Qty",
        "Priced Lines",
        "Unpriced Lines",
        "Total Line Cost",
    ]
    for col_idx, h in enumerate(headers, start=1):
        c = summary.cell(row=5, column=col_idx, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center")
    for i, srow in enumerate(summary_rows, start=6):
        for col_idx, v in enumerate(srow, start=1):
            cell = summary.cell(row=i, column=col_idx, value=v)
            if col_idx == 7:
                cell.number_format = "#,##0.00"
    grand_qty = sum(r[3] for r in summary_rows)
    grand_priced = sum(r[4] for r in summary_rows)
    grand_unpriced = sum(r[5] for r in summary_rows)
    grand_cost = sum(r[6] for r in summary_rows)
    grand_row = 6 + len(summary_rows)
    summary.cell(row=grand_row, column=1, value="TOTAL").font = TOTAL_FONT
    summary.cell(row=grand_row, column=4, value=grand_qty).font = TOTAL_FONT
    summary.cell(row=grand_row, column=5, value=grand_priced).font = TOTAL_FONT
    summary.cell(row=grand_row, column=6, value=grand_unpriced).font = TOTAL_FONT
    gtc = summary.cell(row=grand_row, column=7, value=grand_cost)
    gtc.font = TOTAL_FONT
    gtc.number_format = "#,##0.00"
    for col in range(1, 8):
        summary.cell(row=grand_row, column=col).fill = TOTAL_FILL

    widths = [14, 22, 12, 12, 14, 16, 18]
    for col_idx, w in enumerate(widths, start=1):
        summary.column_dimensions[get_column_letter(col_idx)].width = w

    wb.save(out)
    print(f"Wrote {out.name}")
    print(f"  {len(by_stmt)} statements, {sum(len(v) for v in by_stmt.values())} (stmt, sku) rows")
    print(f"  priced: {grand_priced}, unpriced: {grand_unpriced}")
    print(f"  total qty: {grand_qty:,}, total line cost: ${grand_cost:,.2f}")


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    stem = sys.argv[1]
    src = resolve_source(stem)
    out = DOWNLOADS / f"{stem}_inventory_report.xlsx"
    quarter_label = stem  # e.g. "1-3-2024" → just use the stem

    master_path = MASTER_UPDATED if MASTER_UPDATED.exists() else MASTER_ORIG
    print(f"Source: {src.name}")
    print(f"Master: {master_path.name}")
    master = load_master_prices(master_path)
    print(f"  {len(master)} SKUs, {sum(1 for _, p in master.values() if p)} priced")
    print()
    build_report(src, out, master, quarter_label)


if __name__ == "__main__":
    main()
