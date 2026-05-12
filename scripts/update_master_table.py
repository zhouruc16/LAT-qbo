"""Append SKUs from Q1/Q2/Q3 2024 order detail exports to the Master Table.

Reads Master Table.xlsx (SKU ID, Product name, Price). Reads Order details
from the three quarterly xlsx exports, collects all unique SKUs, and appends
any that are missing from the master with empty Price.

Output: Master Table_updated.xlsx (does not overwrite the original).
"""
from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

DOWNLOADS = Path.home() / "Downloads"
MASTER_SRC = DOWNLOADS / "Master Table.xlsx"
MASTER_OUT = DOWNLOADS / "Master Table_updated.xlsx"

# Quarterly xlsx files. Both Q1 files are processed in case of differences;
# duplicates collapse on (SKU, name) anyway.
QUARTERLY_FILES = [
    DOWNLOADS / "1-3-2024.xlsx",
    DOWNLOADS / "1-3-2024 (1).xlsx",
    DOWNLOADS / "4-6-2024 (1).xlsx",
    DOWNLOADS / "7-9-2024.xlsx",
]


def read_master_skus(path: Path) -> tuple[set[str], list]:
    """Return (existing SKU id set, raw rows list including header)."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    skus: set[str] = set()
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        rows.append(row)
        if i == 1:
            continue
        sku = row[0]
        if sku is None:
            continue
        skus.add(str(sku).strip())
    return skus, rows


def collect_quarterly_skus(files: list[Path]) -> dict[str, str]:
    """Return {sku_id -> most-common product name} across all quarterly files."""
    name_votes: dict[str, Counter] = defaultdict(Counter)
    for path in files:
        if not path.exists():
            print(f"  [skip] {path.name} (not found)")
            continue
        wb = load_workbook(path, data_only=True)
        if "Order details" not in wb.sheetnames:
            print(f"  [skip] {path.name} (no 'Order details' sheet)")
            continue
        ws = wb["Order details"]
        header = [c.value for c in ws[1]]
        try:
            sku_col = header.index("SKU ID")
            name_col = header.index("Product name")
        except ValueError:
            print(f"  [skip] {path.name} (missing SKU/Product columns)")
            continue
        n_rows = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            sku = row[sku_col]
            name = row[name_col]
            if not sku:
                continue
            sku_str = str(sku).strip()
            if not sku_str or sku_str == "None":
                continue
            # Skip non-product placeholder SKUs (TikTok uses "/" for fee-only,
            # adjustment, and reimbursement rows that aren't actual products).
            if not sku_str.isdigit() or len(sku_str) < 15:
                continue
            name_str = str(name).strip() if name else ""
            name_votes[sku_str][name_str] += 1
            n_rows += 1
        print(f"  [read] {path.name}: {n_rows:,} rows with SKU")
    # pick most common name per SKU
    return {sku: counter.most_common(1)[0][0] for sku, counter in name_votes.items()}


def main() -> None:
    print(f"Reading master table: {MASTER_SRC}")
    master_skus, _ = read_master_skus(MASTER_SRC)
    print(f"  master has {len(master_skus)} SKUs")
    print()

    print("Reading quarterly order detail files:")
    quarterly = collect_quarterly_skus(QUARTERLY_FILES)
    print(f"  {len(quarterly)} unique SKUs across all quarterly files")
    print()

    missing = sorted(sku for sku in quarterly if sku not in master_skus)
    print(f"Missing from master table: {len(missing)} SKUs")
    print()

    # write updated workbook
    wb = load_workbook(MASTER_SRC)
    ws = wb.active
    next_row = ws.max_row + 1
    appended_fill = PatternFill("solid", fgColor="FFF2CC")  # pale yellow
    italic = Font(italic=True)

    for sku in missing:
        ws.cell(row=next_row, column=1, value=sku)
        ws.cell(row=next_row, column=2, value=quarterly[sku])
        # column 3 (Price) intentionally left empty
        for col in (1, 2, 3):
            cell = ws.cell(row=next_row, column=col)
            cell.fill = appended_fill
            cell.font = italic
        next_row += 1

    wb.save(MASTER_OUT)
    print(f"Wrote {MASTER_OUT}")
    print(f"  added {len(missing)} rows (highlighted pale yellow, italic)")


if __name__ == "__main__":
    main()
